"""Crée une copie lisible du jeu de données pour Excel : data/processed/corpus_final.xlsx

Usage (depuis la racine du projet) :
    pip install openpyxl
    python src\\collect\\export_excel.py

Ne modifie pas corpus_final.csv. Cette copie sert à lire les données, pas à les partager
avec le code de l'équipe (qui lit le CSV).
"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = Path("data/processed/corpus_final.csv")
OUT = Path("data/processed/corpus_final.xlsx")
WIDTHS = {"article_id": 14, "title": 70, "published_at": 22, "source": 22, "url": 40, "text": 50,
          "text_status": 12, "fetched_at": 22, "theme": 14, "theme_score": 12, "entities": 20,
          "event_id": 10, "sentiment": 11, "is_relevant": 11, "duplicate_of": 14}
EXCEL_MAX = 32000  # une cellule Excel ne dépasse pas 32 767 caractères


def main() -> None:
    csv.field_size_limit(10**9)
    with SRC.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    wb = Workbook()
    ws = wb.active
    ws.title = "articles"
    for r in rows:
        ws.append([c[:EXCEL_MAX] for c in r])
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    for i, cell in enumerate(ws[1], 1):
        cell.font = Font(bold=True)
        cell.fill = head_fill
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(cell.value, 15)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{len(rows) - 1} articles -> {OUT}")


if __name__ == "__main__":
    main()